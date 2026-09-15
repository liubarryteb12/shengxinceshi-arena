#!/usr/bin/env python3
"""Generate a conservative, evidence-anchored P2 L1 manuscript package."""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import shutil
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Cm, Inches, Pt
from docx.oxml.ns import qn


LOCAL_TZ = ZoneInfo("Asia/Shanghai")


def write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.rstrip() + "\n", encoding="utf-8")


def load_summary(root: Path) -> tuple[dict, str]:
    run_dir = root / "analysis/outputs/affymetrix_expression_gse7451/run"
    candidates = sorted(run_dir.glob("summary.json"), key=lambda p: p.stat().st_mtime)
    if not candidates:
        raise SystemExit("P2 requires a P1 summary.json")
    path = candidates[-1]
    summary = json.loads(path.read_text(encoding="utf-8"))
    return summary, summary.get("run_id", path.parent.name)


def setup_doc(document: Document, line_spacing: float) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = line_spacing
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_para(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workspace", type=Path, default=Path("."))
    args = parser.parse_args()
    root = args.workspace.resolve()
    summary, run_id = load_summary(root)
    today = dt.datetime.now(LOCAL_TZ).date().isoformat()
    p2 = root / "02_writing"
    version = "v1"
    version_dir = p2 / "manuscript" / version
    version_dir.mkdir(parents=True, exist_ok=True)
    figures_source_dir = root / "analysis/outputs/affymetrix_expression_gse7451/figures"
    figure_target = p2 / "figures/Figure_1_QC.png"
    for figures_source in sorted(figures_source_dir.glob("Figure_1_QC.*")):
        shutil.copy2(figures_source, p2 / "figures" / figures_source.name)

    n = summary.get("n_samples", 20)
    groups = summary.get("groups", {"control": 10, "pSS": 10})
    filtered = summary.get("filtered_probe_count", "待 P1 full_raw")
    sig = summary.get("significant_gene_count", "待 P1 full_raw")
    input_mode = summary.get("input_mode", "unknown")
    p1_status = "partial" if input_mode == "series_matrix" else "success"
    evidence_note = "本地 matrix_only 验证；raw CEL 的 RMA/MAS5 尚待 GitHub Actions full_raw" if input_mode == "series_matrix" else "raw CEL full_raw 验证"

    title = "Exploratory salivary gene expression analysis in primary Sjögren's syndrome"
    md = f"""> Output level: L1
> P1 status: {p1_status}
> Evidence boundary: candidate / exploratory
> Target journal: pending
> Note: {evidence_note}

# {title}

## One-sentence positioning

This project evaluates whether a reproducible Affymetrix expression workflow can transform GSE7451 salivary expression data into traceable candidate evidence for primary Sjögren's syndrome research.

## Logical chain P0–P4

- **P0 Background problem**: salivary molecular profiles may differ between primary Sjögren's syndrome and healthy controls.
- **P1 Main finding**: the current validation package contains a balanced 10-control/10-pSS design [L-E001].
- **P2 Supporting evidence**: an expression matrix and limma/BH comparison were produced; the run reports {sig} multiple-testing-significant probes under its current input mode [L-E002].
- **P3 Limitation**: this local run is {input_mode}; it does not replace the raw-CEL RMA/MAS5 run.
- **P4 Outlook**: execute the same configuration in GitHub Actions, then review pathway-tool compatibility and independent validation needs.

## Abstract (structured draft)

### Background
Primary Sjögren's syndrome is an autoimmune disease for which minimally invasive molecular signals remain of interest. This work uses a reproducible workflow to test an exploratory salivary expression analysis.

### Methods
The GSE7451 series contains {n} samples, including {groups.get('control', 10)} controls and {groups.get('pSS', 10)} pSS samples [L-E001]. The target pipeline is RMA preprocessing, removal of probe sets absent in more than 75% of samples, limma differential expression with Benjamini–Hochberg correction, and pathway enrichment. The present local run used {input_mode}; the raw-CEL path is scheduled on GitHub Actions.

### Results
The current run produced an expression matrix with {filtered} retained rows and a differential-expression table. {sig} rows met the configured adjusted-P and effect-size thresholds in this run [L-E002]. Figure 1 is a QC visualization, not a clinical-performance plot.

### Conclusions
The repository now records a traceable P1-to-P2 evidence path. The findings are exploratory and candidate-level; independent validation and the raw-CEL execution remain necessary.

## Introduction

Primary Sjögren's syndrome affects exocrine tissues and can be studied with molecular measurements from saliva. The present workflow is designed to separate data processing, evidence extraction, writing, and journal strategy. It therefore treats the GSE7451 result as a candidate signal rather than a validated diagnostic marker.

## Materials and methods

### Dataset and sample groups

GSE7451 is an expression-profiling-by-array series on whole saliva. The downloaded package contains Affymetrix GPL570 raw files and a GEO series matrix. The tracked sample sheet records {n} samples, with balanced control and pSS groups [L-E001].

### Reproducible analysis workflow

The confirmed workflow is: input integrity and sample check → RMA → MAS5 present/absent filtering at the pre-registered 75% threshold → limma comparison of pSS versus control with BH correction → pathway enrichment. Every step is configured through files rather than hard-coded user paths. In the local validation run, the series matrix was used only to validate downstream contracts; the full raw-CEL workflow is assigned to GitHub Actions.

### Evidence and claim boundary

All numerical statements in this draft are anchored to the P1 run and its evidence interface. The terms “candidate”, “exploratory”, “associated”, and “requires independent validation” are intentional. This workflow does not support claims of diagnosis, clinical utility, causality, mechanism confirmation, or external validation.

## Results

### Input and workflow integrity

The input package was downloaded from NCBI GEO, recorded with checksums, and parsed into a sample sheet. The group balance was control={groups.get('control', 10)} and pSS={groups.get('pSS', 10)} [L-E001].

### Differential expression output

The P1 run generated normalized/filtered matrix files and a limma table. The current interface records {filtered} rows after the configured input-mode handling and {sig} threshold-qualified rows [L-E002]. These values are run-specific and must be regenerated after the GitHub Actions raw-CEL run.

[Figure 1 position]

**Figure 1.** GSE7451 expression QC boxplot. The figure is exploratory and does not establish diagnostic performance. Exact sample size and processing mode are defined in the P1 manifest [L-E001].

## Discussion

The main contribution of the current stage is a reproducible handoff rather than a definitive biological conclusion. The interface keeps the raw-CEL requirement, the pathway-tool distinction, and the evidence boundary visible to downstream writing and typesetting. The next technical checkpoint is the GitHub Actions full_raw execution.

### Limitations

The dataset is a single cohort with 10 samples per group, all recorded as female in the GEO sample metadata. There is no independent validation cohort or experimental validation in the downloaded package. The local matrix-only run does not execute raw-CEL RMA or MAS5 calls, and the pathway step is not presented as an original MAPPFinder result.

### Outlook

A future version should compare the full_raw output with the matrix-only validation, activate independent validation if data become available, and update the P4-b journal assessment using verified journal sources.

## Declarations

### Ethics approval and consent to participate
Not applicable to the public GEO reanalysis at this stage; confirm against the target journal and source study.

### Consent for publication
Not applicable; confirm against the target journal.

### Availability of data and materials
The GEO accession is GSE7451. Local file checksums and provenance are recorded in `inputs/migration_log.yaml`.

### Code availability
The workflow scripts are maintained in the `shengxinceshi-arena` repository; remote push is pending GitHub credentials in the execution environment.

### Competing interests
Not assessed; author input required.

### Funding
Not provided in the current project inputs; author input required.

### Authors' contributions
Not provided in the current project inputs; author input required.

### Declaration of generative AI use
The project uses an agent to execute and record workflow steps. The final disclosure must be adapted to the target journal and author policy.

## References

See `02_writing/refs/refs.bib`. References are limited to records present in the supplied workflow specification or GEO metadata.
"""
    write(version_dir / "main.md", md)
    shutil.copy2(version_dir / "main.md", p2 / "manuscript/main.md")

    # Generate editable DOCX versions with the same text and an optional QC image.
    for name, spacing in (("main.docx", 2.0), ("main_审稿版.docx", 2.0), ("main_安全版.docx", 1.5)):
        document = Document()
        setup_doc(document, spacing)
        document.add_heading(title, level=0)
        add_para(document, f"Output level: L1 | Evidence boundary: candidate / exploratory | P1 input mode: {input_mode}")
        for line in md.splitlines():
            if not line or line.startswith(">") or line.startswith("[Figure 1 position]"):
                continue
            if line.startswith("### "):
                document.add_heading(line[4:], level=3)
            elif line.startswith("## "):
                document.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                document.add_heading(line[2:], level=1)
            elif line.startswith("- "):
                document.add_paragraph(line[2:], style="List Bullet")
            elif line.startswith("**Figure 1."):
                add_para(document, line)
            elif line.startswith("See `"):
                add_para(document, line)
            else:
                add_para(document, line)
        if figure_target.exists():
            document.add_picture(str(figure_target), width=Inches(6.0))
            cap = document.add_paragraph("Figure 1. GSE7451 expression QC boxplot; exploratory only.")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        target = version_dir / name
        document.save(target)
    shutil.copy2(version_dir / "main_审稿版.docx", p2 / "manuscript/main_审稿版.docx")
    shutil.copy2(version_dir / "main_安全版.docx", p2 / "manuscript/main_安全版.docx")

    write(p2 / "p2_outline.md", md[:4000])
    write(p2 / "p2_claims.yaml", f'''schema_version: "1.0"\nstatus: partial\nclaims:\n  - claim_id: L-E001\n    text: GSE7451 contains {n} samples with control={groups.get('control', 10)} and pSS={groups.get('pSS', 10)}.\n    evidence: analysis/_index/p1_to_p2_evidence.yaml#E001\n    strength: candidate\n  - claim_id: L-E002\n    text: The current run produced a limma/BH differential-expression table with {sig} threshold-qualified rows.\n    evidence: analysis/_index/p1_to_p2_evidence.yaml#E002\n    strength: exploratory\n''')
    write(p2 / "p2_nc_check.yaml", f'''schema_version: "1.0"\nstatus: partial\nchecks:\n  NC-1: {{verdict: N/A, detail: research question response requires target story confirmation}}\n  NC-2: {{verdict: fail, detail: single-cohort evidence and no independent validation}}\n  NC-3: {{verdict: pass, detail: discussion begins from P1 boundary}}\n  NC-4: {{verdict: pass, detail: result sections are ordered}}\n  NC-5: {{verdict: pass, detail: Figure 1 has a purpose statement}}\n  NC-6: {{verdict: pass, detail: abstract begins from the P1 evidence boundary}}\n  NC-7: {{verdict: pass, detail: no new results are introduced in discussion}}\n''')
    write(p2 / "p2_declarations.yaml", '''schema_version: "1.0"\nstatus: partial\nitems:\n  ethics: Not applicable pending target-journal confirmation\n  consent: Not applicable pending target-journal confirmation\n  data_availability: GEO GSE7451\n  code_availability: repository path recorded; remote push pending credentials\n  competing_interests: author input required\n  funding: author input required\n  authors_contributions: author input required\n  ai_declaration: draft requires target-journal wording\n''')
    write(p2 / "review/review_summary.yaml", f'''schema_version: "1.0"\nreview_id: review_{run_id}_p2\nmanuscript_version: v1\ntimestamp: {today}\nstatus: pending_agent_review\nmodels:\n  - role: domain_expert\n    status: pending\n  - role: methodologist\n    status: pending\n  - role: statistician\n    status: pending\n  - role: journal_editor\n    status: pending\nsummary:\n  total_major: null\n  total_minor: null\n  overall: pending\n''')
    write(p2 / "revision_log.yaml", f'''schema_version: "1.0"\nmanuscript_id: ms_GSE7451_{today.replace('-', '')}\nrevisions:\n  - version: v1\n    timestamp: {today}\n    trigger: initial_l1_draft\n    note: generated from P1 evidence interface; not final submission text\n''')
    write(p2 / "refs/refs.bib", '''@misc{GSE7451,\n  title = {Primary Sjogren's syndrome and control whole saliva},\n  howpublished = {NCBI GEO GSE7451},\n  note = {PMID: 17968930}\n}\n\n@article{limma2015,\n  title = {limma powers differential expression analyses for RNA-sequencing and microarray studies},\n  author = {Ritchie, Matthew E. and Phipson, Belinda and Wu, Di and Hu, Yifang and Law, Charity W. and Shi, Wei and Smyth, Gordon K.},\n  journal = {Nucleic Acids Research},\n  year = {2015},\n  volume = {43},\n  number = {7},\n  pages = {e47},\n  doi = {10.1093/nar/gkv007}\n}\n''')
    write(p2 / "refs/refs_audit.yaml", '''schema_version: "1.0"\nstatus: partial\nsource_policy: only supplied specification and GEO metadata\nentries:\n  - key: GSE7451\n    source: GEO series matrix metadata\n    verified: true\n  - key: limma2015\n    source: docs/source/工作安排提示词.txt\n    verified: true\n''')

    interface = p2 / "p2_to_p3_manuscript.yaml"
    write(interface, f'''schema_version: "1.0"\ncontract_id: p2_to_p3_manuscript\nproject_id: shengxinceshi_arena\nstatus: partial\ntimestamp: {today}\np1_run_id: {run_id}\nmanuscript_version: v1\noutput_level: L1\nstructure:\n  order: S1\n  sections: [title_page, abstract, keywords, introduction, materials_and_methods, results, discussion, declarations, references]\nfiles:\n  manuscript: 02_writing/manuscript/v1/main.docx\n  manuscript_markdown: 02_writing/manuscript/v1/main.md\n  refs: 02_writing/refs/refs.bib\n  figures_dir: 02_writing/figures/\n  tables_dir: 02_writing/tables/\nfigures:\n  - figure_id: Figure_1_QC\n    file: 02_writing/figures/Figure_1_QC.png\n    caption: Figure 1. GSE7451 expression QC boxplot; exploratory only.\n    position_hint: after_results_para_1\ntables:\n  - table_id: Table_DE\n    file: analysis/outputs/affymetrix_expression_gse7451/tables/differential_expression.csv\n    caption: Differential expression results\n    position_hint: after_results_para_1\nreferences:\n  style: Vancouver_pending_target_journal\n  count: 2\ndeclarations:\n  ethics: draft\n  consent: draft\n  data_availability: GEO GSE7451\n  code_availability: draft\n  competing_interests: pending_author_input\n  funding: pending_author_input\n  authors_contributions: pending_author_input\n  ai_declaration: draft\nvalidation:\n  status: partial\n  checks: [WP-0, WP-CORE, WP-4, WP-5, WP-6, WP-7, NC-1~7]\n''')
    write(p2 / "p2_to_p4b_quality.yaml", f'''schema_version: "1.0"\ncontract_id: p2_to_p4b_quality\nproject_id: shengxinceshi_arena\nstatus: partial\ntimestamp: {today}\np1_run_id: {run_id}\nmanuscript_version: v1\nquality_assessment:\n  narrative_completeness: provisional\n  evidence_consistency: provisional\n  method_rigor: provisional\n  statistical_validity: provisional\n  writing_quality: provisional\n  figure_quality: provisional\n  novelty: not_assessed\n  clinical_value: not_assessed\ntopic_assessment:\n  primary_topic: 原发性干燥综合征与健康对照的唾液转录组表达差异\n  secondary_topic: 唾液基因表达候选生物标志物\n  potential_journals: [自身免疫病, 口腔医学, 生物信息学/转录组学]\nstrengths:\n  - 证据锚点已写入稿件\n  - 措辞保守，没有升级为诊断或机制已证实\nweaknesses:\n  - P1 当前为 matrix_only partial run\n  - 缺少独立验证和作者声明信息\nrecommendations:\n  - 等待 full_raw P1 结果后生成 L2\n  - 目标期刊确定后核对结构、摘要、参考文献和声明\n''')
    print(json.dumps({"status": "partial", "output_level": "L1", "manuscript": str(interface.relative_to(root)), "p1_run_id": run_id}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
